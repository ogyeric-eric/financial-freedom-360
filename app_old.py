import io
from datetime import datetime
import numpy as np
import pandas as pd
import streamlit as st

st.set_page_config(page_title='Finansijska sloboda 360', page_icon='🧭', layout='wide')

st.markdown('''
<style>
.block-container {padding-top: 1.1rem; padding-bottom: 2rem; max-width: 1180px;}
[data-testid="stMetric"] {background: rgba(127,127,127,.07); padding: .7rem; border-radius: 12px;}
.smallnote {font-size:.83rem; opacity:.75}
.hero {padding:.9rem 1rem;border-radius:14px;background:rgba(80,120,180,.10);margin-bottom:.7rem}
@media (max-width: 768px) {
  .block-container {padding-left:.8rem;padding-right:.8rem;}
  [data-testid="stMetric"] {padding:.55rem;}
}
</style>
''', unsafe_allow_html=True)

# ---------------- Market data ----------------
@st.cache_data(ttl=1800, show_spinner=False)
def load_market_data(symbols):
    import yfinance as yf
    out = {}
    for key, ticker in symbols.items():
        try:
            h = yf.download(ticker, period='6mo', interval='1d', progress=False, auto_adjust=False, threads=False)
            if h is None or h.empty:
                out[key] = None
                continue
            if isinstance(h.columns, pd.MultiIndex):
                close = h['Close'][ticker] if ticker in h['Close'].columns else h['Close'].iloc[:, 0]
            else:
                close = h['Close']
            close = pd.to_numeric(close, errors='coerce').dropna()
            if len(close) < 22:
                out[key] = None
                continue
            current = float(close.iloc[-1])
            ref_1m = float(close.iloc[-22])
            ref_3m = float(close.iloc[max(0, len(close)-66)])
            ret = np.log(close / close.shift(1)).dropna()
            vol_1m = float(ret.tail(21).std() * np.sqrt(252)) if len(ret) >= 21 else np.nan
            out[key] = {
                'current': current,
                'mom_1m': current/ref_1m - 1 if ref_1m else 0,
                'mom_3m': current/ref_3m - 1 if ref_3m else 0,
                'vol_1m_ann': vol_1m,
                'date': pd.Timestamp(close.index[-1]).date().isoformat(),
            }
        except Exception:
            out[key] = None
    return out

DEFAULT_SYMBOLS = {
    'Gold': 'GC=F',
    'Brent': 'BZ=F',
    'CSPX': 'CSPX.L',
    'Energy': 'ZPDE.DE',
    'VIX': '^VIX',
    'US10Y': '^TNX',
    'EURUSD': 'EURUSD=X',
}

# ---------------- Core logic ----------------
def pyramid_rank(reserve_months, independent_income, min_cost, comfort_cost, target_cost):
    if reserve_months < 6:
        return 1, 'EDUKACIJA'
    if independent_income < min_cost:
        return 2, 'ZAŠTITA'
    if independent_income < comfort_cost:
        return 3, 'SIGURNOST'
    if independent_income < max(target_cost, 1.25 * comfort_cost):
        return 4, 'SLOBODA'
    return 5, 'KOMPETENCIJA'

BASE_WEIGHTS = {
    1: {'Cash': .70, 'Gold': .10, 'CSPX': .15, 'Energy': .05},
    2: {'Cash': .50, 'Gold': .15, 'CSPX': .30, 'Energy': .05},
    3: {'Cash': .30, 'Gold': .15, 'CSPX': .45, 'Energy': .10},
    4: {'Cash': .15, 'Gold': .15, 'CSPX': .55, 'Energy': .15},
    5: {'Cash': .10, 'Gold': .10, 'CSPX': .60, 'Energy': .20},
}
MAX_ENERGY = {1: .05, 2: .05, 3: .10, 4: .15, 5: .20}
OVERLAY = {
    'NEUTRAL': {'Cash': 0, 'Gold': 0, 'CSPX': 0, 'Energy': 0},
    'STRESS': {'Cash': .10, 'Gold': .10, 'CSPX': -.15, 'Energy': -.05},
    'DEFENSIVE': {'Cash': .05, 'Gold': .10, 'CSPX': -.10, 'Energy': -.05},
    'GROWTH': {'Cash': -.05, 'Gold': -.05, 'CSPX': .10, 'Energy': 0},
    'ENERGY / INFLATION': {'Cash': -.05, 'Gold': .05, 'CSPX': -.05, 'Energy': .05},
}

def iri_score(reserve_months, savings_rate, debt_ratio, ifs):
    reserve_score = min(1, max(0, reserve_months / 6))
    savings_score = min(1, max(0, savings_rate) / .25)
    debt_score = 1 if debt_ratio <= .10 else (0 if debt_ratio >= .30 else (.30 - debt_ratio) / .20)
    ifs_score = min(1, max(0, ifs))
    return .35*reserve_score + .25*savings_score + .20*debt_score + .20*ifs_score

def classify_regime(m):
    risk = energy = growth = defensive = 0
    gm, bm, cm, em = m.get('Gold', 0), m.get('Brent', 0), m.get('CSPX', 0), m.get('Energy', 0)
    vix = m.get('VIX', 18)
    yield10 = m.get('US10Y', 4.0)
    defensive += 2 if gm >= .05 else (1 if gm > 0 else 0)
    energy += 3 if bm >= .20 else (2 if bm >= .10 else (1 if bm >= .03 else 0))
    risk += 2 if cm <= -.08 else (1 if cm <= -.03 else 0)
    growth += 3 if cm >= .05 else (2 if cm >= .02 else (1 if cm > 0 else 0))
    defensive += 1 if cm < 0 else 0
    energy += 2 if em >= .08 else (1 if em >= .02 else 0)
    gor_change = gm - bm
    energy += 2 if gor_change <= -.10 else (1 if gor_change <= -.05 else 0)
    defensive += 2 if gor_change >= .05 else (1 if gor_change > 0 else 0)
    risk += 3 if vix >= 30 else (2 if vix >= 25 else (1 if vix >= 20 else 0))
    growth += 2 if vix < 18 else (1 if vix < 22 else 0)
    # Higher yields are a mild headwind for duration/gold; kept low-weight intentionally.
    risk += 1 if yield10 >= 5.0 else 0
    scores = {'STRESS': risk, 'ENERGY / INFLATION': energy, 'GROWTH': growth, 'DEFENSIVE': defensive}
    mx = max(scores.values())
    if mx == 0:
        return 'NEUTRAL', 0.0, scores
    regime = max(scores, key=scores.get)
    return regime, mx / max(1, sum(scores.values())), scores

def target_weights(rank, regime, reserve_months, debt_ratio, ifs):
    base = BASE_WEIGHTS[rank].copy()
    ov = OVERLAY[regime]
    raw = {k: max(0, base[k] + ov[k]) for k in base}
    energy_cap = 0 if reserve_months < 3 else (.05 if reserve_months < 6 or debt_ratio > .30 else (.10 if ifs < 1 else MAX_ENERGY[rank]))
    raw['Energy'] = min(raw['Energy'], energy_cap)
    rem = 1 - raw['Energy']
    subtotal = sum(raw[k] for k in ['Cash', 'Gold', 'CSPX'])
    for k in ['Cash', 'Gold', 'CSPX']:
        raw[k] = raw[k] / subtotal * rem if subtotal else 0
    return raw

def next_level_gap(level, independent, min_cost, comfort, target):
    if level == 'EDUKACIJA':
        return 'ZAŠTITA', None
    if level == 'ZAŠTITA':
        return 'SIGURNOST', max(0, min_cost - independent)
    if level == 'SIGURNOST':
        return 'SLOBODA', max(0, comfort - independent)
    if level == 'SLOBODA':
        goal = max(target, 1.25*comfort)
        return 'KOMPETENCIJA', max(0, goal - independent)
    return 'KOMPETENCIJA', 0

def export_xlsx(summary_df, portfolio_df, alloc_df, market_df):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = 'SAZETAK'
    for r in dataframe_to_rows_safe(summary_df):
        ws.append(r)
    for name, df in [('PORTFELJ', portfolio_df), ('NOVI_KAPITAL', alloc_df), ('TRZISTE', market_df)]:
        s = wb.create_sheet(name)
        for r in dataframe_to_rows_safe(df):
            s.append(r)
    for s in wb.worksheets:
        s.freeze_panes = 'A2'
        for cell in s[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F4E78')
            cell.alignment = Alignment(horizontal='center')
        for col in s.columns:
            width = min(35, max(11, max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2))
            s.column_dimensions[col[0].column_letter].width = width
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.getvalue()

def dataframe_to_rows_safe(df):
    yield list(df.columns)
    for _, row in df.iterrows():
        yield [None if pd.isna(v) else v for v in row.tolist()]

# ---------------- Inputs ----------------
st.title('🧭 Finansijska sloboda 360')
st.markdown('<div class="hero"><b>Jednostavan cilj:</b> gdje sam danas → šta tržište govori → šta mi nedostaje u portfelju → šta sa sljedećih X KM → koliko sam bliže slobodi.</div>', unsafe_allow_html=True)

with st.sidebar:
    st.header('1) Finansijska baza')
    min_cost = st.number_input('Minimalni troškovi / mj. (KM)', min_value=0.0, value=0.0, step=100.0)
    comfort = st.number_input('Komforni troškovi / mj. (KM)', min_value=0.0, value=0.0, step=100.0)
    target = st.number_input('Ciljni standard / mj. (KM)', min_value=0.0, value=0.0, step=100.0)
    actual_spend = st.number_input('Stvarna potrošnja / mj. (KM)', min_value=0.0, value=0.0, step=100.0)
    net_income = st.number_input('Ukupan neto prihod / mj. (KM)', min_value=0.0, value=0.0, step=100.0)
    independent = st.number_input('Nezavisni prihod / mj. (KM)', min_value=0.0, value=0.0, step=100.0)
    monthly_save = st.number_input('Mjesečno ulaganje/štednja (KM)', min_value=0.0, value=0.0, step=100.0)
    debt_service = st.number_input('Mjesečna otplata duga (KM)', min_value=0.0, value=0.0, step=50.0)

    st.header('2) Imovina')
    total_cash = st.number_input('Gotovina i depoziti (KM)', min_value=0.0, value=0.0, step=500.0)
    reserve = st.number_input('Od toga: zaštitna rezerva (KM)', min_value=0.0, value=0.0, step=500.0)
    physical_gold = st.number_input('Zlatne kovanice / poluge (KM)', min_value=0.0, value=0.0, step=500.0)
    gold_etf = st.number_input('Gold ETF / drugo zlato (KM)', min_value=0.0, value=0.0, step=500.0)
    cspx = st.number_input('S&P 500 UCITS / CSPX (KM)', min_value=0.0, value=0.0, step=500.0)
    energy_etf = st.number_input('Energy UCITS (KM)', min_value=0.0, value=0.0, step=500.0)
    bonds = st.number_input('Obveznice / money market (KM)', min_value=0.0, value=0.0, step=500.0)
    other_fin = st.number_input('Ostali ETF / akcije / fondovi (KM)', min_value=0.0, value=0.0, step=500.0)
    other_assets = st.number_input('Nekretnine / poslovni udjeli / drugo (KM)', min_value=0.0, value=0.0, step=1000.0)
    new_capital = st.number_input('Sljedeći kapital za raspodjelu (KM)', min_value=0.0, value=0.0, step=500.0)

reserve_months = reserve / actual_spend if actual_spend else 0
ifs = independent / comfort if comfort else 0
savings_rate = monthly_save / net_income if net_income else 0
debt_ratio = debt_service / net_income if net_income else 0
rank, level = pyramid_rank(reserve_months, independent, min_cost, comfort, target)
iri = iri_score(reserve_months, savings_rate, debt_ratio, ifs)

# ---------------- Market ----------------
with st.expander('⚙️ Tržišni feed i tickeri', expanded=False):
    st.caption('Tickeri su podesivi. Ako feed zakaže, aplikacija automatski nudi ručni fallback.')
    symbols = {}
    cols = st.columns(2)
    items = list(DEFAULT_SYMBOLS.items())
    for i, (k, v) in enumerate(items):
        symbols[k] = cols[i % 2].text_input(k, value=v, key=f'ticker_{k}')

market = load_market_data(symbols)
required = ['Gold', 'Brent', 'CSPX', 'Energy', 'VIX']
manual = not all(market.get(k) for k in required)
if manual:
    with st.expander('⚠️ Ručni fallback tržišta', expanded=True):
        st.warning('Online feed nije kompletan. Unesite približne 1M promjene; ostali dijelovi aplikacije nastavljaju raditi.')
        c1, c2 = st.columns(2)
        gm = c1.number_input('Gold 1M %', value=0.0) / 100
        bm = c1.number_input('Brent 1M %', value=0.0) / 100
        cm = c2.number_input('CSPX 1M %', value=0.0) / 100
        em = c2.number_input('Energy 1M %', value=0.0) / 100
        vix = c2.number_input('VIX', value=18.0)
        us10y = c1.number_input('US 10Y %', value=4.0)
        mom = {'Gold': gm, 'Brent': bm, 'CSPX': cm, 'Energy': em, 'VIX': vix, 'US10Y': us10y}
else:
    mom = {k: (market[k]['mom_1m'] if k not in ['VIX','US10Y'] else market[k]['current']) for k in required + ['US10Y'] if market.get(k)}

regime, confidence, scores = classify_regime(mom)
targets = target_weights(rank, regime, reserve_months, debt_ratio, ifs)

# ---------------- Data frames ----------------
invest_cash = max(0, total_cash - reserve)
current = {'Cash': invest_cash, 'Gold': physical_gold + gold_etf, 'CSPX': cspx, 'Energy': energy_etf}
active_total = sum(current.values())
portfolio_rows = []
for key, label in [('Cash','Investabilna gotovina'), ('Gold','Zlato ukupno'), ('CSPX','S&P 500 UCITS'), ('Energy','Energy UCITS')]:
    actual = current[key]/active_total if active_total else 0
    tgt = targets[key]
    gap = tgt*active_total - current[key]
    threshold = max(100, .01*active_total)
    action = 'POVEĆAJ' if gap > threshold else ('NE POVEĆAVAJ / REBALANS' if gap < -threshold else 'ZADRŽI')
    portfolio_rows.append([label, current[key], actual, tgt, gap, action])
portfolio_df = pd.DataFrame(portfolio_rows, columns=['Aktiva','Trenutno KM','Trenutno %','Cilj %','Gap KM','Akcija'])

reserve_gap = max(0, actual_spend*6 - reserve)
reserve_alloc = min(new_capital, reserve_gap)
remaining = max(0, new_capital - reserve_alloc)
future_total = active_total + remaining
positive_gaps = {k: max(0, targets[k]*future_total - current[k]) for k in current}
gsum = sum(positive_gaps.values())
alloc = {k: (remaining*(positive_gaps[k]/gsum) if gsum else remaining*targets[k]) for k in current}
alloc_df = pd.DataFrame([
    ['Dopuna rezerve', reserve_alloc], ['Investabilna gotovina', alloc['Cash']], ['Gold', alloc['Gold']],
    ['CSPX', alloc['CSPX']], ['Energy UCITS', alloc['Energy']]
], columns=['Namjena','Predloženo KM'])

if not manual:
    mrows = []
    for k in DEFAULT_SYMBOLS:
        d = market.get(k)
        if not d: continue
        mrows.append([k, symbols[k], d['current'], d['mom_1m'] if k not in ['VIX','US10Y','EURUSD'] else np.nan, d['date']])
    market_df = pd.DataFrame(mrows, columns=['Instrument','Ticker','Zadnja vrijednost','1M momentum','Datum'])
else:
    market_df = pd.DataFrame([[k, symbols.get(k,''), v, np.nan, 'manual'] for k,v in mom.items()], columns=['Instrument','Ticker','Zadnja vrijednost','1M momentum','Datum'])

# ---------------- Tabs ----------------
tab_today, tab_portfolio, tab_new, tab_path, tab_comp = st.tabs(['🏠 DANAS','💼 MOJ PORTFELJ','➕ NOVI KAPITAL','🛤️ PUT DO SLOBODE','🎯 KOMPETENCIJA'])

with tab_today:
    c1,c2,c3,c4,c5 = st.columns(5)
    c1.metric('Nivo', level)
    c2.metric('IFS', f'{ifs:.0%}')
    c3.metric('IRI', f'{iri:.0%}')
    c4.metric('Rezerva', f'{reserve_months:.1f} mj.')
    c5.metric('Tržišni režim', regime)
    if iri >= .80: readiness = 'SPREMAN ZA DINAMIČKU ALOKACIJU'
    elif iri >= .60: readiness = 'USLOVNO SPREMAN — OGRANIČEN RIZIK'
    else: readiness = 'PRVO OJAČATI FINANSIJSKU BAZU'
    st.info(f'**Investiciona spremnost:** {readiness}. Pouzdanost režimskog signala: {confidence:.0%}.')
    if reserve_months < 6:
        st.warning('Prioritet je dopuna zaštitne rezerve prije povećanja rizičnijeg dijela portfelja.')
    elif debt_ratio > .30:
        st.warning('Servis duga prelazi 30% neto prihoda — ograničiti satellite rizik i analizirati smanjenje duga.')
    elif regime == 'ENERGY / INFLATION':
        st.success('Tržište favorizuje energetsku komponentu, ali ona ostaje satellite i podliježe cap-u prema Vašem nivou.')
    elif regime == 'GROWTH':
        st.success('Growth režim daje relativnu prednost širokom core equity dijelu (S&P 500 UCITS).')
    elif regime in ['STRESS','DEFENSIVE']:
        st.success('Defanzivni režim: veći značaj likvidnosti i zaštite; novi rizik uvoditi postepeno.')
    else:
        st.success('Neutralan režim: slijediti ciljnu piramidu i portfolio gap, bez agresivnog taktičkog pomjeranja.')
    st.subheader('Tržišni pregled')
    st.dataframe(market_df, use_container_width=True, hide_index=True)

with tab_portfolio:
    total_net_assets = total_cash + physical_gold + gold_etf + cspx + energy_etf + bonds + other_fin + other_assets
    c1,c2,c3 = st.columns(3)
    c1.metric('Ukupna evidentirana aktiva', f'{total_net_assets:,.0f} KM')
    c2.metric('Aktivni investicioni portfolio', f'{active_total:,.0f} KM')
    c3.metric('Fizičko zlato', f'{physical_gold:,.0f} KM')
    st.subheader('Portfolio gap')
    styled = portfolio_df.style.format({'Trenutno KM':'{:,.0f}','Trenutno %':'{:.1%}','Cilj %':'{:.1%}','Gap KM':'{:,.0f}'})
    st.dataframe(styled, use_container_width=True, hide_index=True)
    st.caption('Nekretnine, obveznice i ostala aktiva ulaze u pregled ukupnog bogatstva, dok taktički gap trenutno upravlja četiri osnovna sloja: Cash, Gold, CSPX i Energy.')

with tab_new:
    st.subheader(f'Šta sa sljedećih {new_capital:,.0f} KM?')
    cols = st.columns(5)
    values = [('Rezerva', reserve_alloc), ('Gotovina', alloc['Cash']), ('Gold', alloc['Gold']), ('CSPX', alloc['CSPX']), ('Energy', alloc['Energy'])]
    for c,(name,val) in zip(cols,values): c.metric(name, f'{val:,.0f} KM')
    st.dataframe(alloc_df.style.format({'Predloženo KM':'{:,.0f}'}), use_container_width=True, hide_index=True)
    if reserve_alloc > 0:
        st.warning(f'Prvo dopuniti rezervu za približno {reserve_alloc:,.0f} KM. Tek ostatak raspoređivati po investicionom gapu.')
    else:
        st.info('Rezerva je na minimalnom cilju; novi kapital se raspoređuje prema portfolio gapu i tržišnom režimu.')

with tab_path:
    next_level, monthly_gap = next_level_gap(level, independent, min_cost, comfort, target)
    st.subheader(f'Sljedeći nivo: {next_level}')
    if level == 'EDUKACIJA':
        gap_reserve = max(0, actual_spend*6 - reserve)
        months = gap_reserve/monthly_save if monthly_save > 0 else np.nan
        st.metric('Nedostaje do 6M rezerve', f'{gap_reserve:,.0f} KM')
        if np.isfinite(months): st.metric('Procjena tempom štednje', f'{months:.1f} mj.')
    elif monthly_gap is not None:
        st.metric('Nedostaje nezavisnog prihoda', f'{monthly_gap:,.0f} KM/mj.')
        capital_equiv = monthly_gap*12/.04 if monthly_gap > 0 else 0
        st.metric('Kapital-ekvivalent pri 4% godišnje', f'{capital_equiv:,.0f} KM')
        st.caption('Kapital-ekvivalent je ilustrativan planerski pokazatelj, ne garancija prinosa niti preporučena stopa povlačenja.')
    st.subheader('Financial Freedom Accelerator')
    st.write('Prioritet se bira po tome šta trenutno najviše skraćuje put do sljedećeg nivoa: rezerva, smanjenje skupog duga, povećanje mjesečnog cash-flowa ili produktivni kapital.')
    priorities=[]
    if reserve_months < 6: priorities.append(('1', 'Dopuniti rezervu', 'Visok'))
    if debt_ratio > .20: priorities.append((str(len(priorities)+1), 'Analizirati ubrzanu otplatu duga', 'Visok'))
    if savings_rate < .20: priorities.append((str(len(priorities)+1), 'Povećati stopu štednje / prihod', 'Srednje-visok'))
    priorities.append((str(len(priorities)+1), 'Automatizovati mjesečno ulaganje u core portfolio', 'Kontinuirano'))
    st.dataframe(pd.DataFrame(priorities, columns=['Red','Akcija','Prioritet']), use_container_width=True, hide_index=True)

with tab_comp:
    diversification = min(1, sum(1 for v in [current['Cash'],current['Gold'],current['CSPX'],current['Energy']] if v > 0)/4)
    discipline = min(1, savings_rate/.20) if savings_rate > 0 else 0
    capital_eff = min(1, active_total / max(1, total_net_assets)) if total_net_assets else 0
    competence = .30*min(1,ifs/1.25) + .25*iri + .20*diversification + .15*discipline + .10*capital_eff
    c1,c2,c3 = st.columns(3)
    c1.metric('Competence Index', f'{competence:.0%}')
    c2.metric('Diversifikacija', f'{diversification:.0%}')
    c3.metric('Disciplina ulaganja', f'{discipline:.0%}')
    st.progress(float(min(1, competence)))
    st.write('Kompetencija ovdje znači: finansijska baza + sistem ulaganja + diversifikacija + disciplina + efikasno korištenje kapitala. Nije samo visina kapitala.')
    summary_df = pd.DataFrame([
        ['Nivo', level], ['IFS', ifs], ['IRI', iri], ['Rezerva mjeseci', reserve_months], ['Tržišni režim', regime], ['Competence Index', competence], ['Vrijeme izvještaja', datetime.now().isoformat(timespec='seconds')]
    ], columns=['Pokazatelj','Vrijednost'])
    xlsx = export_xlsx(summary_df, portfolio_df, alloc_df, market_df)
    st.download_button('⬇️ Izvezi trenutni izvještaj u Excel', data=xlsx, file_name='Financial_Freedom_360_snapshot.xlsx', mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

st.divider()
st.caption('Decision-support alat za planiranje i disciplinu. Tržišni podaci mogu kasniti; prije svake stvarne transakcije provjeriti cijenu, poreze, troškove i dostupnost instrumenta kod brokera.')
