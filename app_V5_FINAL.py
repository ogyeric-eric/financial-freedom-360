
import io
from datetime import datetime

import numpy as np
import pandas as pd
import streamlit as st

st.set_page_config(page_title="Finansijska sloboda 360", page_icon="🧭", layout="wide")

st.markdown("""
<style>
.block-container {padding-top: 1.1rem; padding-bottom: 2rem; max-width: 1180px;}
[data-testid="stMetric"] {background: rgba(127,127,127,.07); padding: .7rem; border-radius: 12px;}
.hero {padding:.9rem 1rem;border-radius:14px;background:rgba(80,120,180,.10);margin-bottom:.7rem}
@media (max-width: 768px) {
  .block-container {padding-left:.8rem;padding-right:.8rem;}
  [data-testid="stMetric"] {padding:.55rem;}
}
</style>
""", unsafe_allow_html=True)

# ---------------- Market data ----------------
@st.cache_data(ttl=1800, show_spinner=False)
def load_market_data(symbols):
    import yfinance as yf
    out = {}
    for key, ticker in symbols.items():
        try:
            h = yf.download(
                ticker, period="6mo", interval="1d",
                progress=False, auto_adjust=False, threads=False
            )
            if h is None or h.empty:
                out[key] = None
                continue

            if isinstance(h.columns, pd.MultiIndex):
                if "Close" in h.columns.get_level_values(0):
                    close_block = h["Close"]
                    if hasattr(close_block, "columns"):
                        close = close_block[ticker] if ticker in close_block.columns else close_block.iloc[:, 0]
                    else:
                        close = close_block
                else:
                    out[key] = None
                    continue
            else:
                close = h["Close"]

            close = pd.to_numeric(close, errors="coerce").dropna()
            if len(close) < 22:
                out[key] = None
                continue

            current = float(close.iloc[-1])
            ref_1m = float(close.iloc[-22])
            ref_3m = float(close.iloc[max(0, len(close) - 66)])
            ret = np.log(close / close.shift(1)).dropna()

            out[key] = {
                "current": current,
                "mom_1m": current / ref_1m - 1 if ref_1m else 0.0,
                "mom_3m": current / ref_3m - 1 if ref_3m else 0.0,
                "vol_1m_ann": float(ret.tail(21).std() * np.sqrt(252)) if len(ret) >= 21 else np.nan,
                "date": pd.Timestamp(close.index[-1]).date().isoformat(),
            }
        except Exception:
            out[key] = None
    return out


DEFAULT_SYMBOLS = {
    "Gold": "GC=F",
    "Brent": "BZ=F",
    "CSPX": "CSPX.L",
    "Energy": "ZPDE.DE",
    "VIX": "^VIX",
    "US10Y": "^TNX",
    "EURUSD": "EURUSD=X",
}

# ---------------- Core logic ----------------
BASE_WEIGHTS = {
    1: {"Cash": .70, "Gold": .10, "CSPX": .15, "Energy": .05},
    2: {"Cash": .50, "Gold": .15, "CSPX": .30, "Energy": .05},
    3: {"Cash": .30, "Gold": .15, "CSPX": .45, "Energy": .10},
    4: {"Cash": .15, "Gold": .15, "CSPX": .55, "Energy": .15},
    5: {"Cash": .10, "Gold": .10, "CSPX": .60, "Energy": .20},
}
MAX_ENERGY = {1: .05, 2: .05, 3: .10, 4: .15, 5: .20}
OVERLAY = {
    "NEUTRAL": {"Cash": 0, "Gold": 0, "CSPX": 0, "Energy": 0},
    "STRESS": {"Cash": .10, "Gold": .10, "CSPX": -.15, "Energy": -.05},
    "DEFENSIVE": {"Cash": .05, "Gold": .10, "CSPX": -.10, "Energy": -.05},
    "GROWTH": {"Cash": -.05, "Gold": -.05, "CSPX": .10, "Energy": 0},
    "ENERGY / INFLATION": {"Cash": -.05, "Gold": .05, "CSPX": -.05, "Energy": .05},
}

def pyramid_rank(reserve_months, independent_income, min_cost, comfort_cost, target_cost):
    if reserve_months < 6:
        return 1, "EDUKACIJA"
    if independent_income < min_cost:
        return 2, "ZAŠTITA"
    if independent_income < comfort_cost:
        return 3, "SIGURNOST"
    if independent_income < max(target_cost, 1.25 * comfort_cost):
        return 4, "SLOBODA"
    return 5, "KOMPETENCIJA"

def iri_score(reserve_months, savings_rate, debt_ratio, ifs):
    reserve_score = min(1, max(0, reserve_months / 6))
    savings_score = min(1, max(0, savings_rate) / .25)
    debt_score = 1 if debt_ratio <= .10 else (0 if debt_ratio >= .30 else (.30 - debt_ratio) / .20)
    ifs_score = min(1, max(0, ifs))
    return .35 * reserve_score + .25 * savings_score + .20 * debt_score + .20 * ifs_score

def classify_regime(m):
    risk = energy = growth = defensive = 0
    gm = m.get("Gold", 0)
    bm = m.get("Brent", 0)
    cm = m.get("CSPX", 0)
    em = m.get("Energy", 0)
    vix = m.get("VIX", 18)
    yield10 = m.get("US10Y", 4.0)

    defensive += 2 if gm >= .05 else (1 if gm > 0 else 0)
    energy += 3 if bm >= .20 else (2 if bm >= .10 else (1 if bm >= .03 else 0))
    risk += 2 if cm <= -.08 else (1 if cm <= -.03 else 0)
    growth += 3 if cm >= .05 else (2 if cm >= .02 else (1 if cm > 0 else 0))
    defensive += 1 if cm < 0 else 0
    energy += 2 if em >= .08 else (1 if em >= .02 else 0)

    gor_change = gm - bm
    energy += 2 if gor_change <= -.10 else (1 if gor_change <= -.05 else 0)
    defensive += 2 if gor_change >= .05 else (1 if gor_change > 0 else 0)

    risk += 3 if vix >= 30 else (2 if vix >= 25 else (1 if vix >= 20 else 0))
    growth += 2 if vix < 18 else (1 if vix < 22 else 0)
    risk += 1 if yield10 >= 5.0 else 0

    scores = {
        "STRESS": risk,
        "ENERGY / INFLATION": energy,
        "GROWTH": growth,
        "DEFENSIVE": defensive,
    }
    mx = max(scores.values())
    if mx == 0:
        return "NEUTRAL", 0.0, scores
    regime = max(scores, key=scores.get)
    return regime, mx / max(1, sum(scores.values())), scores

def target_weights(rank, regime, reserve_months, debt_ratio, ifs):
    base = BASE_WEIGHTS[rank].copy()
    overlay = OVERLAY[regime]
    raw = {k: max(0, base[k] + overlay[k]) for k in base}

    energy_cap = (
        0 if reserve_months < 3
        else .05 if reserve_months < 6 or debt_ratio > .30
        else .10 if ifs < 1
        else MAX_ENERGY[rank]
    )
    raw["Energy"] = min(raw["Energy"], energy_cap)

    rem = 1 - raw["Energy"]
    subtotal = sum(raw[k] for k in ["Cash", "Gold", "CSPX"])
    for k in ["Cash", "Gold", "CSPX"]:
        raw[k] = raw[k] / subtotal * rem if subtotal else 0
    return raw

def next_level_gap(level, independent, min_cost, comfort, target):
    if level == "EDUKACIJA":
        return "ZAŠTITA", None
    if level == "ZAŠTITA":
        return "SIGURNOST", max(0, min_cost - independent)
    if level == "SIGURNOST":
        return "SLOBODA", max(0, comfort - independent)
    if level == "SLOBODA":
        return "KOMPETENCIJA", max(0, max(target, 1.25 * comfort) - independent)
    return "KOMPETENCIJA", 0

def months_to_target(pv, pmt, annual_rate, fv, max_months=1200):
    if fv <= pv:
        return 0
    r = annual_rate / 12
    if pmt <= 0 and r <= 0:
        return None
    for n in range(1, max_months + 1):
        if r > 0:
            value = pv * ((1 + r) ** n) + pmt * (((1 + r) ** n - 1) / r)
        else:
            value = pv + pmt * n
        if value >= fv:
            return n
    return None

def dataframe_rows(df):
    yield list(df.columns)
    for _, row in df.iterrows():
        yield [None if pd.isna(v) else v for v in row.tolist()]

def export_xlsx(summary_df, portfolio_df, alloc_df, market_df):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "SAZETAK"
    for r in dataframe_rows(summary_df):
        ws.append(r)

    for name, df in [
        ("PORTFELJ", portfolio_df),
        ("NOVI_KAPITAL", alloc_df),
        ("TRZISTE", market_df),
    ]:
        s = wb.create_sheet(name)
        for r in dataframe_rows(df):
            s.append(r)

    for s in wb.worksheets:
        s.freeze_panes = "A2"
        for cell in s[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
        for col in s.columns:
            width = min(35, max(11, max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2))
            s.column_dimensions[col[0].column_letter].width = width

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.getvalue()

# ---------------- Header ----------------
st.title("🧭 PERSONAL 360 — V5")
st.markdown(
    '<div class="hero"><b>Od finansijske baze do kompetencije:</b> '
    'finansije → vrijednosti → ciljevi → prioriteti → obrasci → kompetencije.</div>',
    unsafe_allow_html=True,
)
st.caption(
    "🔒 Privatnost: uneseni iznosi koriste se u aktivnoj Streamlit sesiji; "
    "ova verzija ih ne zapisuje u GitHub niti u vlastitu bazu podataka."
)

# ---------------- Inputs ----------------
with st.sidebar:
    st.header("1) Finansijska baza")
    min_cost = st.number_input("Minimalni troškovi / mj. (KM)", min_value=0.0, value=0.0, step=100.0)
    comfort = st.number_input("Komforni troškovi / mj. (KM)", min_value=0.0, value=0.0, step=100.0)
    target = st.number_input("Ciljni standard / mj. (KM)", min_value=0.0, value=0.0, step=100.0)
    actual_spend = st.number_input("Stvarna potrošnja / mj. (KM)", min_value=0.0, value=0.0, step=100.0)
    net_income = st.number_input("Ukupan neto prihod / mj. (KM)", min_value=0.0, value=0.0, step=100.0)
    independent = st.number_input("Nezavisni prihod / mj. (KM)", min_value=0.0, value=0.0, step=100.0)
    monthly_save = st.number_input("Mjesečno ulaganje/štednja (KM)", min_value=0.0, value=0.0, step=100.0)
    debt_service = st.number_input("Mjesečna otplata duga (KM)", min_value=0.0, value=0.0, step=50.0)
    debt_balance = st.number_input("Preostali dug ukupno (KM)", min_value=0.0, value=0.0, step=500.0)
    debt_rate = st.number_input(
        "Prosječna kamata na dug (% godišnje)",
        min_value=0.0, max_value=50.0, value=0.0, step=0.1
    ) / 100

    st.header("2) Imovina")
    total_cash = st.number_input("Gotovina i depoziti (KM)", min_value=0.0, value=0.0, step=500.0)
    reserve = st.number_input("Od toga: zaštitna rezerva (KM)", min_value=0.0, value=0.0, step=500.0)
    physical_gold = st.number_input("Zlatne kovanice / poluge (KM)", min_value=0.0, value=0.0, step=500.0)
    gold_etf = st.number_input("Gold ETF / drugo zlato (KM)", min_value=0.0, value=0.0, step=500.0)
    cspx = st.number_input("S&P 500 UCITS / CSPX (KM)", min_value=0.0, value=0.0, step=500.0)
    energy_etf = st.number_input("Energy UCITS (KM)", min_value=0.0, value=0.0, step=500.0)
    bonds = st.number_input("Obveznice / money market (KM)", min_value=0.0, value=0.0, step=500.0)
    other_fin = st.number_input("Ostali ETF / akcije / fondovi (KM)", min_value=0.0, value=0.0, step=500.0)
    other_assets = st.number_input("Nekretnine / poslovni udjeli / drugo (KM)", min_value=0.0, value=0.0, step=1000.0)
    new_capital = st.number_input("Sljedeći kapital za raspodjelu (KM)", min_value=0.0, value=0.0, step=500.0)

reserve_months = reserve / actual_spend if actual_spend else 0
ifs = independent / comfort if comfort else 0
savings_rate = monthly_save / net_income if net_income else 0
debt_ratio = debt_service / net_income if net_income else 0
rank, level = pyramid_rank(reserve_months, independent, min_cost, comfort, target)
iri = iri_score(reserve_months, savings_rate, debt_ratio, ifs)

# ---------------- Market ----------------
with st.expander("⚙️ Tržišni feed i tickeri", expanded=False):
    st.caption("Tickeri su podesivi. Ako feed zakaže, aplikacija automatski nudi ručni fallback.")
    symbols = {}
    cols = st.columns(2)
    for i, (k, v) in enumerate(DEFAULT_SYMBOLS.items()):
        symbols[k] = cols[i % 2].text_input(k, value=v, key=f"ticker_{k}")

market = load_market_data(symbols)
required = ["Gold", "Brent", "CSPX", "Energy", "VIX"]
manual = not all(market.get(k) for k in required)

if manual:
    with st.expander("⚠️ Ručni fallback tržišta", expanded=True):
        st.warning("Online feed nije kompletan. Unesite približne 1M promjene.")
        c1, c2 = st.columns(2)
        gm = c1.number_input("Gold 1M %", value=0.0) / 100
        bm = c1.number_input("Brent 1M %", value=0.0) / 100
        cm = c2.number_input("CSPX 1M %", value=0.0) / 100
        em = c2.number_input("Energy 1M %", value=0.0) / 100
        vix = c2.number_input("VIX", value=18.0)
        us10y = c1.number_input("US 10Y %", value=4.0)
        mom = {"Gold": gm, "Brent": bm, "CSPX": cm, "Energy": em, "VIX": vix, "US10Y": us10y}
else:
    mom = {
        k: (market[k]["mom_1m"] if k not in ["VIX", "US10Y"] else market[k]["current"])
        for k in required + ["US10Y"]
        if market.get(k)
    }

regime, confidence, scores = classify_regime(mom)
targets = target_weights(rank, regime, reserve_months, debt_ratio, ifs)

# ---------------- Portfolio ----------------
invest_cash = max(0, total_cash - reserve)
current = {
    "Cash": invest_cash,
    "Gold": physical_gold + gold_etf,
    "CSPX": cspx,
    "Energy": energy_etf,
}
active_total = sum(current.values())

portfolio_rows = []
for key, label in [
    ("Cash", "Investabilna gotovina"),
    ("Gold", "Zlato ukupno"),
    ("CSPX", "S&P 500 UCITS"),
    ("Energy", "Energy UCITS"),
]:
    actual = current[key] / active_total if active_total else 0
    tgt = targets[key]
    gap = tgt * active_total - current[key]
    threshold = max(100, .01 * active_total)
    action = "POVEĆAJ" if gap > threshold else ("NE POVEĆAVAJ / REBALANS" if gap < -threshold else "ZADRŽI")
    portfolio_rows.append([label, current[key], actual, tgt, gap, action])

portfolio_df = pd.DataFrame(
    portfolio_rows,
    columns=["Aktiva", "Trenutno KM", "Trenutno %", "Cilj %", "Gap KM", "Akcija"],
)

reserve_gap = max(0, actual_spend * 6 - reserve)
reserve_alloc = min(new_capital, reserve_gap)
remaining = max(0, new_capital - reserve_alloc)
future_total = active_total + remaining
positive_gaps = {k: max(0, targets[k] * future_total - current[k]) for k in current}
gsum = sum(positive_gaps.values())
alloc = {
    k: (remaining * positive_gaps[k] / gsum if gsum else remaining * targets[k])
    for k in current
}
alloc_df = pd.DataFrame(
    [
        ["Dopuna rezerve", reserve_alloc],
        ["Investabilna gotovina", alloc["Cash"]],
        ["Gold", alloc["Gold"]],
        ["CSPX", alloc["CSPX"]],
        ["Energy UCITS", alloc["Energy"]],
    ],
    columns=["Namjena", "Predloženo KM"],
)

if not manual:
    market_rows = []
    for k in DEFAULT_SYMBOLS:
        d = market.get(k)
        if not d:
            continue
        market_rows.append(
            [
                k, symbols[k], d["current"],
                d["mom_1m"] if k not in ["VIX", "US10Y", "EURUSD"] else np.nan,
                d["date"],
            ]
        )
    market_df = pd.DataFrame(
        market_rows,
        columns=["Instrument", "Ticker", "Zadnja vrijednost", "1M momentum", "Datum"],
    )
else:
    market_df = pd.DataFrame(
        [[k, symbols.get(k, ""), v, np.nan, "manual"] for k, v in mom.items()],
        columns=["Instrument", "Ticker", "Zadnja vrijednost", "1M momentum", "Datum"],
    )

total_assets = total_cash + physical_gold + gold_etf + cspx + energy_etf + bonds + other_fin + other_assets


# ---------------- Personal 360 state ----------------
if "personal_values" not in st.session_state:
    st.session_state.personal_values = []
if "personal_goals" not in st.session_state:
    st.session_state.personal_goals = []
if "eisenhower_tasks" not in st.session_state:
    st.session_state.eisenhower_tasks = []

def value_names():
    return [v["Naziv"] for v in st.session_state.personal_values]

def goal_names():
    return [g["Naziv"] for g in st.session_state.personal_goals]

# ---------------- Tabs ----------------
tabs = st.tabs([
    "🏠 DANAS",
    "💼 MOJ PORTFELJ",
    "➕ NOVI KAPITAL",
    "🛤️ PUT DO SLOBODE",
    "🚀 ACCELERATOR",
    "❤️ VRIJEDNOSTI",
    "🎯 CILJEVI",
    "⏱️ EISENHOWER",
    "🎓 KOMPETENCIJA",
])

with tabs[0]:
    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Nivo", level)
    c2.metric("IFS", f"{ifs:.0%}")
    c3.metric("IRI", f"{iri:.0%}")
    c4.metric("Rezerva", f"{reserve_months:.1f} mj.")
    c5.metric("Tržišni režim", regime)

    readiness = (
        "SPREMAN ZA DINAMIČKU ALOKACIJU"
        if iri >= .80
        else "USLOVNO SPREMAN — OGRANIČEN RIZIK"
        if iri >= .60
        else "PRVO OJAČATI FINANSIJSKU BAZU"
    )
    st.info(f"**Investiciona spremnost:** {readiness}. Pouzdanost režimskog signala: {confidence:.0%}.")

    if reserve_months < 6:
        st.warning("Prioritet je dopuna zaštitne rezerve prije povećanja rizičnijeg dijela portfelja.")
    elif debt_ratio > .30:
        st.warning("Servis duga prelazi 30% neto prihoda — ograničiti satellite rizik i analizirati smanjenje duga.")
    elif regime == "ENERGY / INFLATION":
        st.success("Tržište favorizuje energetsku komponentu, ali ona ostaje satellite i podliježe limitu piramide.")
    elif regime == "GROWTH":
        st.success("Growth režim daje relativnu prednost širokom core equity dijelu (S&P 500 UCITS).")
    elif regime in ["STRESS", "DEFENSIVE"]:
        st.success("Defanzivni režim: veći značaj likvidnosti i zaštite; novi rizik uvoditi postepeno.")
    else:
        st.success("Neutralan režim: slijediti ciljnu piramidu i portfolio gap bez agresivnog taktičkog pomjeranja.")

    st.subheader("Tržišni pregled")
    st.dataframe(market_df, use_container_width=True, hide_index=True)

with tabs[1]:
    c1, c2, c3 = st.columns(3)
    c1.metric("Ukupna evidentirana aktiva", f"{total_assets:,.0f} KM")
    c2.metric("Aktivni investicioni portfolio", f"{active_total:,.0f} KM")
    c3.metric("Fizičko zlato", f"{physical_gold:,.0f} KM")
    st.subheader("Portfolio gap")
    st.dataframe(
        portfolio_df.style.format({
            "Trenutno KM": "{:,.0f}",
            "Trenutno %": "{:.1%}",
            "Cilj %": "{:.1%}",
            "Gap KM": "{:,.0f}",
        }),
        use_container_width=True,
        hide_index=True,
    )

with tabs[2]:
    st.subheader(f"Šta sa sljedećih {new_capital:,.0f} KM?")
    cols = st.columns(5)
    values = [
        ("Rezerva", reserve_alloc),
        ("Gotovina", alloc["Cash"]),
        ("Gold", alloc["Gold"]),
        ("CSPX", alloc["CSPX"]),
        ("Energy", alloc["Energy"]),
    ]
    for c, (name, val) in zip(cols, values):
        c.metric(name, f"{val:,.0f} KM")
    st.dataframe(
        alloc_df.style.format({"Predloženo KM": "{:,.0f}"}),
        use_container_width=True,
        hide_index=True,
    )

with tabs[3]:
    next_level, monthly_gap = next_level_gap(level, independent, min_cost, comfort, target)
    st.subheader(f"Sljedeći nivo: {next_level}")

    if level == "EDUKACIJA":
        gap_reserve = max(0, actual_spend * 6 - reserve)
        months = gap_reserve / monthly_save if monthly_save > 0 else np.nan
        st.metric("Nedostaje do 6M rezerve", f"{gap_reserve:,.0f} KM")
        if np.isfinite(months):
            st.metric("Procjena tempom štednje", f"{months:.1f} mj.")
    elif monthly_gap is not None:
        st.metric("Nedostaje nezavisnog prihoda", f"{monthly_gap:,.0f} KM/mj.")
        capital_equiv = monthly_gap * 12 / .04 if monthly_gap > 0 else 0
        st.metric("Kapital-ekvivalent pri 4% godišnje", f"{capital_equiv:,.0f} KM")
        st.caption("Kapital-ekvivalent je ilustrativan planerski pokazatelj, ne garancija prinosa.")

with tabs[4]:
    st.subheader("Financial Freedom Accelerator")

    a1, a2, a3 = st.columns(3)
    extra_income = a1.number_input(
        "Dodatni neto prihod / mj. (KM)", min_value=0.0, value=0.0, step=100.0
    )
    expense_cut = a2.number_input(
        "Trajno smanjenje rashoda / mj. (KM)", min_value=0.0, value=0.0, step=100.0
    )
    expected_return = a3.number_input(
        "Planerski prinos portfelja (% godišnje)",
        min_value=0.0, max_value=15.0, value=5.0, step=0.5
    ) / 100

    accelerated_capacity = max(0, monthly_save + extra_income + expense_cut)
    st.metric(
        "Novi mjesečni investicioni kapacitet",
        f"{accelerated_capacity:,.0f} KM",
        delta=f"{extra_income + expense_cut:,.0f} KM/mj.",
    )

    if debt_balance > 0 and debt_rate > 0:
        annual_interest = debt_balance * debt_rate
        st.info(
            f"Procijenjeni godišnji trošak kamate na prijavljeni dug: "
            f"{annual_interest:,.0f} KM. To je koristan prag pri poređenju "
            f"ubrzane otplate duga i novog ulaganja."
        )

    if level == "ZAŠTITA":
        target_income_for_level = min_cost
    elif level == "SIGURNOST":
        target_income_for_level = comfort
    else:
        target_income_for_level = max(target, 1.25 * comfort)

    goal_capital = (
        max(0, target_income_for_level * 12 / .04)
        if level not in ["EDUKACIJA", "KOMPETENCIJA"]
        else 0
    )
    productive_capital = active_total + bonds + other_fin

    base_months = (
        months_to_target(productive_capital, monthly_save, expected_return, goal_capital)
        if goal_capital else None
    )
    accelerated_months = (
        months_to_target(productive_capital, accelerated_capacity, expected_return, goal_capital)
        if goal_capital else None
    )

    x1, x2, x3, x4 = st.columns(4)
    x1.metric("Produktivni kapital", f"{productive_capital:,.0f} KM")
    x2.metric("Planerski ciljni kapital", f"{goal_capital:,.0f} KM" if goal_capital else "—")
    x3.metric("Bazni horizont", f"{base_months} mj." if base_months is not None else "—")
    if base_months is not None and accelerated_months is not None:
        x4.metric("Moguće ubrzanje", f"{base_months - accelerated_months} mj.")
    else:
        x4.metric("Moguće ubrzanje", "—")

    priorities = []
    if reserve_months < 6:
        priorities.append(["Dopuniti rezervu do najmanje 6 mjeseci", "VISOK"])
    if debt_ratio > .20 or debt_rate >= .07:
        priorities.append(["Analizirati ubrzanu otplatu skupljeg duga", "VISOK"])
    if savings_rate < .20:
        priorities.append(["Povećati mjesečni višak: prihod + kontrola rashoda", "SREDNJE-VISOK"])
    priorities.append(["Automatizovati mjesečno ulaganje u core portfolio", "KONTINUIRANO"])
    priorities.append(["Satellite pozicije držati unutar limita piramide", "KONTROLISANO"])

    st.dataframe(
        pd.DataFrame(priorities, columns=["Akcija", "Prioritet"]),
        use_container_width=True,
        hide_index=True,
    )
    st.caption(
        "Planerski prinos, 4% kapital-ekvivalent i procjena vremena su scenarijske pretpostavke, "
        "ne garancija budućeg prinosa."
    )

with tabs[5]:
    st.subheader("❤️ Lične vrijednosti")
    st.caption("Definišite 5–10 vrijednosti, rangirajte ih i procijenite koliko ih trenutno živite.")
    with st.form("value_form", clear_on_submit=True):
        v1,v2,v3=st.columns([2,1,1])
        value_name=v1.text_input("Naziv vrijednosti", placeholder="npr. Porodica, Sloboda, Integritet, Znanje")
        value_importance=v2.slider("Važnost",1,10,8)
        value_lived=v3.slider("Koliko je živim danas",1,10,5)
        value_desc=st.text_area("Šta ova vrijednost konkretno znači u mom životu?",height=80)
        add_value=st.form_submit_button("➕ Dodaj vrijednost",use_container_width=True)
    if add_value:
        if value_name.strip():
            st.session_state.personal_values.append({"Naziv":value_name.strip(),"Važnost":int(value_importance),"Živim_danas":int(value_lived),"Gap":int(value_importance-value_lived),"Opis":value_desc.strip()})
            st.success("Vrijednost je dodata.")
        else: st.warning("Unesite naziv vrijednosti.")
    if st.session_state.personal_values:
        values_df=pd.DataFrame(st.session_state.personal_values).sort_values(by=["Važnost","Gap"],ascending=[False,False]).reset_index(drop=True)
        values_df.insert(0,"Rang",range(1,len(values_df)+1))
        st.dataframe(values_df[["Rang","Naziv","Važnost","Živim_danas","Gap","Opis"]],use_container_width=True,hide_index=True)
        align=values_df["Živim_danas"].sum()/values_df["Važnost"].sum() if values_df["Važnost"].sum()>0 else 0
        top=values_df.iloc[0]["Naziv"]; gaprow=values_df.sort_values("Gap",ascending=False).iloc[0]
        a,b,c=st.columns(3);a.metric("Vrijednost #1",top);b.metric("Indeks usklađenosti",f"{align:.0%}");c.metric("Najveći gap",f"{gaprow['Naziv']} ({gaprow['Gap']})")
        st.caption("Indeks usklađenosti je opisni pokazatelj između deklarisane važnosti i procjene koliko vrijednost trenutno živite.")
    else:
        st.info("Dodajte svoje ključne vrijednosti. One će postati osnova za ciljeve i prioritete.")

with tabs[6]:
    st.subheader("🎯 Ciljevi")
    st.caption("Svaki cilj može biti povezan sa oblastima života i Vašim ličnim vrijednostima.")
    value_options=value_names()
    with st.form("goal_form", clear_on_submit=True):
        g1,g2=st.columns([2,1])
        goal_name=g1.text_input("Naziv cilja",placeholder="npr. Dostići finansijsku slobodu")
        goal_area=g2.selectbox("Oblast",["Finansije","Posao/Karijera","Porodica","Zdravlje","Lični razvoj","Naučni rad","Strani jezici","Ostalo"])
        g3,g4,g5=st.columns(3)
        goal_horizon=g3.selectbox("Horizont",["30 dana","90 dana","1 godina","3 godine","5+ godina"])
        goal_priority=g4.slider("Prioritet",1,10,8)
        goal_progress=g5.slider("Napredak %",0,100,0)
        linked_values=st.multiselect("Poveži sa vrijednostima",value_options,default=[]) if value_options else []
        goal_why=st.text_area("Zašto mi je ovaj cilj važan?",height=70)
        goal_measure=st.text_input("Kako ću znati da je cilj ostvaren?",placeholder="mjerljiv kriterij / rezultat")
        add_goal=st.form_submit_button("➕ Dodaj cilj",use_container_width=True)
    if add_goal:
        if goal_name.strip():
            st.session_state.personal_goals.append({"Naziv":goal_name.strip(),"Oblast":goal_area,"Horizont":goal_horizon,"Prioritet":int(goal_priority),"Napredak":int(goal_progress),"Vrijednosti":", ".join(linked_values),"Zašto":goal_why.strip(),"Mjerilo":goal_measure.strip(),"Status":"AKTIVAN"})
            st.success("Cilj je dodat.")
        else: st.warning("Unesite naziv cilja.")
    if st.session_state.personal_goals:
        goals_df=pd.DataFrame(st.session_state.personal_goals)
        active=goals_df[goals_df["Status"]=="AKTIVAN"].copy()
        st.dataframe(active[["Naziv","Oblast","Horizont","Prioritet","Napredak","Vrijednosti","Mjerilo"]],use_container_width=True,hide_index=True)
        weighted=((active["Napredak"]*active["Prioritet"]).sum()/active["Prioritet"].sum()) if not active.empty and active["Prioritet"].sum()>0 else 0
        a,b,c=st.columns(3);a.metric("Aktivnih ciljeva",len(active));b.metric("Ponderisani napredak",f"{weighted:.0f}%");c.metric("Vezanih za vrijednosti",int((active["Vrijednosti"].str.len()>0).sum()) if not active.empty else 0)
        choice=st.selectbox("Ažuriraj cilj",["—"]+[g["Naziv"] for g in st.session_state.personal_goals],key="goal_update_select")
        if choice!="—":
            selected=next(g for g in st.session_state.personal_goals if g["Naziv"]==choice)
            newp=st.slider("Novi napredak %",0,100,int(selected["Napredak"]),key="goal_progress_update")
            u1,u2=st.columns(2)
            if u1.button("Sačuvaj napredak",use_container_width=True):
                selected["Napredak"]=int(newp); selected["Status"]="OSTVAREN" if newp>=100 else "AKTIVAN"; st.rerun()
            if u2.button("Označi cilj kao ostvaren",use_container_width=True):
                selected["Napredak"]=100;selected["Status"]="OSTVAREN";st.rerun()
    else:
        st.info("Dodajte 3–7 aktivnih ciljeva. Kasnije ćemo pratiti njihovu usklađenost sa dnevnim aktivnostima.")

with tabs[7]:
    st.subheader("⏱️ Eisenhower matrica")
    st.caption("BITNO/HITNO · BITNO/NIJE HITNO · NIJE BITNO/HITNO · NIJE BITNO/NIJE HITNO. Zadatak se može direktno povezati sa već unesenim ciljem i ličnom vrijednošću.")

    if "eisenhower_tasks" not in st.session_state:
        st.session_state.eisenhower_tasks = []

    with st.form("eisenhower_form", clear_on_submit=True):
        c1, c2 = st.columns([2,1])
        task_name = c1.text_input("Zadatak / aktivnost")
        area = c2.selectbox("Oblast", ["Posao/Karijera","Finansije","Lični razvoj","Zdravlje","Porodica","Naučni rad","Strani jezici","Ostalo"])
        c3, c4, c5 = st.columns(3)
        important = c3.selectbox("Bitnost", ["BITNO","NIJE BITNO"])
        urgent = c4.selectbox("Hitnost", ["HITNO","NIJE HITNO"])
        deadline = c5.date_input("Rok", value=None)
        c6, c7 = st.columns(2)
        current_goals = ["— Bez veze —"] + goal_names()
        current_values = ["— Bez veze —"] + value_names()
        goal_choice = c6.selectbox("Poveži sa ciljem", current_goals)
        value_choice = c7.selectbox("Poveži sa vrijednošću", current_values)
        goal_link = "" if goal_choice == "— Bez veze —" else goal_choice
        value_link = "" if value_choice == "— Bez veze —" else value_choice
        note = st.text_area("Napomena", height=70)
        add_task = st.form_submit_button("➕ Dodaj u matricu", use_container_width=True)

    if add_task:
        if task_name.strip():
            if important == "BITNO" and urgent == "HITNO": q, action = "I — BITNO / HITNO", "URADI"
            elif important == "BITNO": q, action = "II — BITNO / NIJE HITNO", "PLANIRAJ"
            elif urgent == "HITNO": q, action = "III — NIJE BITNO / HITNO", "DELEGIRAJ"
            else: q, action = "IV — NIJE BITNO / NIJE HITNO", "ELIMINIŠI / OGRANIČI"
            st.session_state.eisenhower_tasks.append({"Zadatak":task_name.strip(),"Oblast":area,"Kvadrant":q,"Akcija":action,"Rok":deadline.isoformat() if deadline else "","Cilj":goal_link.strip(),"Vrijednost":value_link.strip(),"Status":"AKTIVNO","Napomena":note.strip()})
            st.success(f"Zadatak je svrstan u: {q} → {action}.")
        else:
            st.warning("Unesite naziv zadatka.")

    tasks = st.session_state.eisenhower_tasks
    st.markdown("#### Moja 4 kvadranta")
    boxes = st.columns(2) + st.columns(2)
    specs = [("I — BITNO / HITNO","URADI","I —"),("II — BITNO / NIJE HITNO","PLANIRAJ","II —"),("III — NIJE BITNO / HITNO","DELEGIRAJ","III —"),("IV — NIJE BITNO / NIJE HITNO","ELIMINIŠI / OGRANIČI","IV —")]
    for box,(title,action,prefix) in zip(boxes,specs):
        with box:
            st.markdown(f"##### {title}"); st.caption(f"Akcija: **{action}**")
            subset=[t for t in tasks if t["Kvadrant"].startswith(prefix) and t["Status"]=="AKTIVNO"]
            if subset:
                for t in subset: st.write(f"• **{t['Zadatak']}** — {t['Oblast']}" + (f" · rok {t['Rok']}" if t['Rok'] else ""))
            else: st.caption("Nema aktivnih zadataka.")

    if tasks:
        st.markdown("#### Evidencija")
        df = pd.DataFrame(tasks)
        st.dataframe(df[["Zadatak","Oblast","Kvadrant","Akcija","Rok","Cilj","Vrijednost","Status"]], use_container_width=True, hide_index=True)
        active=[(i,t) for i,t in enumerate(tasks) if t["Status"]=="AKTIVNO"]
        if active:
            labels=[f"{i+1}. {t['Zadatak']}" for i,t in active]
            selected=st.selectbox("Izaberite zadatak za završavanje",labels)
            if st.button("✓ Označi kao završeno", use_container_width=True):
                idx=int(selected.split('.',1)[0])-1; st.session_state.eisenhower_tasks[idx]["Status"]="ZAVRŠENO"; st.rerun()
        counts=[sum(1 for t in tasks if t["Kvadrant"].startswith(p) and t["Status"]=="AKTIVNO") for p in ["I —","II —","III —","IV —"]]
        m=st.columns(4)
        for col,label,val in zip(m,["I Bitno/Hitno","II Bitno/Nije hitno","III Nije bitno/Hitno","IV Nije bitno/Nije hitno"],counts): col.metric(label,val)
        important_total=counts[0]+counts[1]; strategic=counts[1]/important_total if important_total else 0
        st.metric("Strateški udio među bitnim zadacima", f"{strategic:.0%}")
        st.download_button("⬇️ Izvezi Eisenhower evidenciju (CSV)", df.to_csv(index=False).encode('utf-8-sig'), "Personal_360_Eisenhower.csv", "text/csv", use_container_width=True)
    else:
        st.info("Dodajte prvi zadatak. Matrica će se automatski popuniti.")

    st.divider()
    st.markdown("#### Prostor za naredne module")
    st.write("Vrijednosti, ciljevi i Eisenhower zadaci sada su povezani. Arhitektura ostaje otvorena za **obrasce/trigere, navike, Active Questions i zajednički Personal 360 dashboard**.")

with tabs[8]:
    diversification = min(1, sum(1 for v in current.values() if v > 0) / 4)
    discipline = min(1, savings_rate / .20) if savings_rate > 0 else 0
    capital_eff = min(1, active_total / max(1, total_assets)) if total_assets else 0
    competence = (
        .30 * min(1, ifs / 1.25)
        + .25 * iri
        + .20 * diversification
        + .15 * discipline
        + .10 * capital_eff
    )

    c1, c2, c3 = st.columns(3)
    c1.metric("Competence Index", f"{competence:.0%}")
    c2.metric("Diversifikacija", f"{diversification:.0%}")
    c3.metric("Disciplina ulaganja", f"{discipline:.0%}")
    st.progress(float(min(1, competence)))

    checks = pd.DataFrame(
        [
            ["Rezerva ≥ 6 mjeseci", reserve_months >= 6],
            ["Servis duga ≤ 20% prihoda", debt_ratio <= .20],
            ["Stopa štednje/ulaganja ≥ 20%", savings_rate >= .20],
            ["Core S&P 500 UCITS postoji", cspx > 0],
            ["Zaštitna komponenta postoji", (physical_gold + gold_etf + reserve) > 0],
            ["Energy ostaje satellite", (energy_etf <= cspx) if cspx > 0 else energy_etf == 0],
        ],
        columns=["Kriterij", "Ispunjeno"],
    )
    checks["Status"] = checks["Ispunjeno"].map({True: "✓", False: "—"})
    st.dataframe(checks[["Kriterij", "Status"]], use_container_width=True, hide_index=True)

    summary_df = pd.DataFrame(
        [
            ["Nivo", level],
            ["IFS", ifs],
            ["IRI", iri],
            ["Rezerva mjeseci", reserve_months],
            ["Tržišni režim", regime],
            ["Competence Index", competence],
            ["Vrijeme izvještaja", datetime.now().isoformat(timespec="seconds")],
        ],
        columns=["Pokazatelj", "Vrijednost"],
    )
    xlsx = export_xlsx(summary_df, portfolio_df, alloc_df, market_df)
    st.download_button(
        "⬇️ Izvezi trenutni izvještaj u Excel",
        data=xlsx,
        file_name="Personal_360_V5_snapshot.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

st.divider()
st.caption(
    "Decision-support alat za planiranje i disciplinu. Tržišni podaci mogu kasniti; "
    "prije stvarne transakcije provjerite cijenu, poreze, troškove, likvidnost i dostupnost instrumenta kod brokera."
)
